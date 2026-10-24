"""Result export: TSV / CSV (optionally gzipped), formatted multi-sheet Excel, Parquet, JSON, and an HTML report."""
from __future__ import annotations

import base64
import datetime as _dt
import html
import json
import os
from typing import Dict, Iterable, List, Optional, Sequence

import numpy as np
import pandas as pd

P_LIKE = ("pval", "qval", "p_", "_p", "fdr", "bh", "by", "bonferroni", "holm", "sidak", "lfdr", "PP.")


def _is_p(col: str) -> bool:
    c = col.lower()
    return any(k.lower() in c for k in P_LIKE)


def write_table(df: pd.DataFrame, stem: str, formats: Sequence[str] = ("tsv",), float_precision: int = 6) -> List[str]:
    """Write one table in each requested format: tsv, tsv.gz, csv, csv.gz, parquet, json, xlsx."""
    os.makedirs(os.path.dirname(stem) or ".", exist_ok=True)
    paths = []
    for f in formats:
        f = f.lower()
        path = f"{stem}.{f}"
        if f.startswith("tsv"):
            df.to_csv(path, sep="\t", index=False, float_format=f"%.{float_precision}g")
        elif f.startswith("csv"):
            df.to_csv(path, index=False, float_format=f"%.{float_precision}g")
        elif f == "parquet":
            try:
                df.to_parquet(path, index=False)
            except Exception as e:  # pyarrow / fastparquet missing
                print(f"[multiqtl] parquet skipped for {stem}: {e.__class__.__name__}")
                continue
        elif f == "json":
            df.to_json(path, orient="records", double_precision=10)
        elif f == "xlsx":
            write_excel({os.path.basename(stem)[:31]: df}, path)
        else:
            raise ValueError(f"unknown format {f}")
        paths.append(path)
    return paths


def write_excel(sheets: Dict[str, pd.DataFrame], path: str, max_rows: int = 1_000_000, readme: Optional[Dict[str, str]] = None):
    """Formatted workbook: bold frozen header, autofilter, column widths, scientific format for p/q-value columns,
    optional README sheet describing every table."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        if readme:
            pd.DataFrame({"sheet": list(readme), "description": list(readme.values())}).to_excel(xw, sheet_name="README", index=False)
        for name, df in sheets.items():
            name = name[:31]
            d = df.head(max_rows).copy()
            for c in d.columns:
                if not pd.api.types.is_numeric_dtype(d[c]) and not pd.api.types.is_bool_dtype(d[c]):
                    d[c] = d[c].astype(str).str.slice(0, 32000)
            d.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            if d.shape[1]:
                ws.auto_filter.ref = ws.dimensions
            for j, c in enumerate(d.columns, 1):
                cell = ws.cell(1, j)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="256ABF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                w = max(len(str(c)), *(len(f"{v:.4g}") if isinstance(v, float) else len(str(v)) for v in d[c].head(200))) if len(d) else len(str(c))
                ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 8), 45)
                if pd.api.types.is_float_dtype(d[c]):
                    fmt = "0.00E+00" if _is_p(str(c)) else "0.0000"
                    for row in ws.iter_rows(min_row=2, min_col=j, max_col=j, max_row=len(d) + 1):
                        row[0].number_format = fmt
        if readme:
            ws = xw.sheets["README"]
            ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 110
            for c in ws[1]:
                c.font = Font(bold=True)
    return path


def to_jsonable(obj):
    if isinstance(obj, dict):
        return {str(k): to_jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [to_jsonable(v) for v in obj]
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (np.floating,)):
        return None if np.isnan(obj) else float(obj)
    if isinstance(obj, float) and np.isnan(obj):
        return None
    if isinstance(obj, (np.bool_,)):
        return bool(obj)
    if isinstance(obj, np.ndarray):
        return to_jsonable(obj.tolist())
    return obj


def write_json(obj, path: str):
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(to_jsonable(obj), fh, indent=2)
    return path


# ================================================================================================================
# HTML report
# ================================================================================================================
_CSS = """
:root{--ink:#0b0b0b;--ink2:#52514e;--muted:#8a8984;--line:#e4e3df;--surface:#fcfcfb;--page:#f9f9f7;--accent:#256abf}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){--ink:#f0efec;--ink2:#c3c2b7;--muted:#8f8e86;--line:#383835;--surface:#1a1a19;--page:#0d0d0d;--accent:#6da7ec}}
:root[data-theme="dark"]{--ink:#f0efec;--ink2:#c3c2b7;--muted:#8f8e86;--line:#383835;--surface:#1a1a19;--page:#0d0d0d;--accent:#6da7ec}
*{box-sizing:border-box}body{margin:0;background:var(--page);color:var(--ink);font:15px/1.55 -apple-system,"Segoe UI",Inter,Helvetica,Arial,sans-serif}
main{max-width:1100px;margin:0 auto;padding:32px 16px 64px}h1{font-size:28px;margin:0 0 4px;letter-spacing:-.01em}
h2{font-size:19px;margin:40px 0 10px;padding-top:12px;border-top:1px solid var(--line)}h3{font-size:15px;margin:22px 0 6px}
.sub{color:var(--ink2);margin:0 0 24px}.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:10px;margin:18px 0}
.kpi{background:var(--surface);border:1px solid var(--line);border-radius:10px;padding:12px 14px}.kpi b{display:block;font-size:24px;font-variant-numeric:tabular-nums}
.kpi span{color:var(--ink2);font-size:12.5px}figure{margin:14px 0 22px;background:#fff;border:1px solid var(--line);border-radius:10px;padding:10px}
figure img{width:100%;height:auto;display:block}figcaption{color:#52514e;font-size:12.5px;margin-top:6px}
.grid2{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:14px}
.tw{overflow-x:auto;border:1px solid var(--line);border-radius:10px;background:var(--surface)}
table{border-collapse:collapse;width:100%;font-size:12.5px;font-variant-numeric:tabular-nums}th,td{padding:6px 9px;border-bottom:1px solid var(--line);text-align:right;white-space:nowrap}
th{position:sticky;top:0;background:var(--surface);color:var(--ink2);font-weight:600}td:first-child,th:first-child{text-align:left}
code{font-size:12.5px}.note{color:var(--ink2);font-size:13px}ul{padding-left:20px}
"""


def _fmt(v):
    if isinstance(v, (float, np.floating)):
        if np.isnan(v):
            return "–"
        if v != 0 and (abs(v) < 1e-3 or abs(v) >= 1e5):
            return f"{v:.2e}"
        return f"{v:.3g}"
    return html.escape(str(v))


def df_to_html(df: pd.DataFrame, max_rows: int = 25) -> str:
    d = df.head(max_rows)
    head = "".join(f"<th>{html.escape(str(c))}</th>" for c in d.columns)
    body = "".join("<tr>" + "".join(f"<td>{_fmt(v)}</td>" for v in r) + "</tr>" for r in d.itertuples(index=False))
    more = f'<p class="note">Showing {len(d)} of {len(df)} rows – full table in the results folder.</p>' if len(df) > max_rows else ""
    return f'<div class="tw"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>{more}'


def img_tag(path: str, max_px: int = 1400) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".svg":
        with open(path, "rb") as fh:
            data = base64.b64encode(fh.read()).decode()
        return f'<img alt="" src="data:image/svg+xml;base64,{data}">'
    # downscale high-dpi rasters for the report (full-resolution files stay in figures/)
    try:
        from io import BytesIO
        from PIL import Image
        im = Image.open(path)
        if im.width > max_px:
            im = im.resize((max_px, int(im.height * max_px / im.width)), Image.LANCZOS)
        im = im.convert("RGB").quantize(colors=128, method=Image.Quantize.MEDIANCUT)   # charts: few colours
        buf = BytesIO()
        im.save(buf, format="PNG", optimize=True)
        data = base64.b64encode(buf.getvalue()).decode()
    except Exception:
        with open(path, "rb") as fh:
            data = base64.b64encode(fh.read()).decode()
    return f'<img alt="" src="data:image/png;base64,{data}">'


class HTMLReport:
    """Self-contained HTML report (figures embedded) assembled section by section."""

    def __init__(self, title: str, subtitle: str = ""):
        self.title, self.subtitle, self.parts = title, subtitle, []

    def h2(self, t):
        self.parts.append(f"<h2>{html.escape(t)}</h2>")

    def h3(self, t):
        self.parts.append(f"<h3>{html.escape(t)}</h3>")

    def p(self, t, raw=False):
        self.parts.append(f"<p>{t if raw else html.escape(t)}</p>")

    def kpis(self, items: Dict[str, object]):
        cells = "".join(f'<div class="kpi"><b>{_fmt(v)}</b><span>{html.escape(k)}</span></div>' for k, v in items.items())
        self.parts.append(f'<div class="kpis">{cells}</div>')

    def figure(self, path: str, caption: str = ""):
        self.parts.append(f"<figure>{img_tag(path)}<figcaption>{html.escape(caption)}</figcaption></figure>")

    def figures(self, items: Iterable[tuple]):
        inner = "".join(f"<figure>{img_tag(p)}<figcaption>{html.escape(c)}</figcaption></figure>" for p, c in items)
        self.parts.append(f'<div class="grid2">{inner}</div>')

    def table(self, df: pd.DataFrame, max_rows: int = 25):
        self.parts.append(df_to_html(df, max_rows))

    def list(self, items: Iterable[str]):
        self.parts.append("<ul>" + "".join(f"<li>{i}</li>" for i in items) + "</ul>")

    def save(self, path: str):
        stamp = _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
        doc = (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
               f"<title>{html.escape(self.title)}</title><style>{_CSS}</style></head><body><main>"
               f"<h1>{html.escape(self.title)}</h1><p class='sub'>{html.escape(self.subtitle)} · generated {stamp}</p>"
               + "".join(self.parts) + "</main></body></html>")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(doc)
        return path
